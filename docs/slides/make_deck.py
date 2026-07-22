# Generates the AI-workshop deck as pptx, mirroring the HTML version:
# dark log-styled slides, JSONL "record" eyebrow, chunk-file footer.
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG = RGBColor(0x10, 0x16, 0x23)
RAISE = RGBColor(0x16, 0x1F, 0x30)
INK = RGBColor(0xE9, 0xE4, 0xD6)
DIM = RGBColor(0x8B, 0x95, 0xA7)
LINE = RGBColor(0x26, 0x32, 0x4A)
INFO = RGBColor(0x7F, 0xB4, 0xC9)
WARN = RGBColor(0xE8, 0xA3, 0x3D)
ERROR = RGBColor(0xD9, 0x6A, 0x5B)
OK = RGBColor(0x97, 0xB8, 0x77)
MONO = "Menlo"
SANS = "Helvetica Neue"

LVL_COLOR = {"info": INFO, "warn": WARN, "error": ERROR, "ok": OK,
             "case": OK, "session": OK, "session_end": OK}

W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def slide_base(seq, lvl, path):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # eyebrow: JSONL record
    tb = s.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    for text, color in [
        ('{"seq":%d,"lvl":' % seq, DIM),
        ('"%s"' % lvl, LVL_COLOR.get(lvl, INFO)),
        (',"path":"%s"}' % path, DIM),
    ]:
        r = p.add_run(); r.text = text
        r.font.name = MONO; r.font.size = Pt(12); r.font.color.rgb = color
    # footer: chunk file name + counter
    fb = s.shapes.add_textbox(Inches(0.7), Inches(7.02), Inches(12), Inches(0.35))
    fp = fb.text_frame.paragraphs[0]
    r = fp.add_run()
    r.text = "tlog_20260723_workshop_p%02d.jsonl" % seq
    r.font.name = MONO; r.font.size = Pt(10); r.font.color.rgb = DIM
    r2 = fp.add_run(); r2.text = "   %d / 18" % seq
    r2.font.name = MONO; r2.font.size = Pt(10); r2.font.color.rgb = LINE
    return s


def title(s, text, size=40, top=0.95, color=INK, accent_words=()):
    tb = s.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(1.1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    for word in text.split(" "):
        r = p.add_run(); r.text = word + " "
        r.font.name = MONO; r.font.size = Pt(size); r.font.bold = True
        r.font.color.rgb = WARN if word.strip(",.") in accent_words else color
    return tb


def bullets(s, items, top=2.1, left=0.7, width=12.0, size=17, gap=8):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(4.6))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for head, rest in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = "> "
        r.font.name = MONO; r.font.size = Pt(size); r.font.color.rgb = WARN
        if head:
            r = p.add_run(); r.text = head + " "
            r.font.name = SANS; r.font.size = Pt(size); r.font.bold = True
            r.font.color.rgb = INK
        if rest:
            r = p.add_run(); r.text = rest
            r.font.name = SANS; r.font.size = Pt(size); r.font.color.rgb = DIM
    return tb


def paragraphs(s, texts, top=2.1, size=18, width=11.8):
    tb = s.shapes.add_textbox(Inches(0.7), Inches(top), Inches(width), Inches(4.6))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for t in texts:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(12)
        r = p.add_run(); r.text = t
        r.font.name = SANS; r.font.size = Pt(size); r.font.color.rgb = INK
        p.line_spacing = 1.25
    return tb


def code_block(s, lines, top=2.05, left=0.7, width=11.9, size=13, height=None):
    height = height or Inches(0.34 * len(lines) + 0.5)
    box = s.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), height)
    box.fill.solid(); box.fill.fore_color.rgb = RAISE
    box.line.color.rgb = LINE; box.line.width = Pt(0.75)
    tf = box.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.18)
    first = True
    for text, color in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = text
        r.font.name = MONO; r.font.size = Pt(size)
        r.font.color.rgb = color
    return box


def table(s, rows, top=2.0, col_w=(3.2, 8.7), size=14, header=("", "")):
    n = len(rows) + (1 if header[0] else 0)
    shape = s.shapes.add_table(n, 2, Inches(0.7), Inches(top),
                               Inches(col_w[0] + col_w[1]), Inches(0.5 * n))
    t = shape.table
    t.columns[0].width = Inches(col_w[0]); t.columns[1].width = Inches(col_w[1])
    data = ([header] if header[0] else []) + rows
    for i, (a, b) in enumerate(data):
        for j, text in enumerate((a, b)):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RAISE if i % 2 == 0 else BG
            cell.vertical_anchor = MSO_ANCHOR.TOP
            cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.07); cell.margin_bottom = Inches(0.07)
            p = cell.text_frame.paragraphs[0]
            r = p.add_run(); r.text = text
            is_header = header[0] and i == 0
            r.font.size = Pt(size - (1 if is_header else 0))
            if j == 0:
                r.font.name = MONO
                r.font.color.rgb = DIM if is_header else INFO
            else:
                r.font.name = SANS
                r.font.color.rgb = DIM if is_header else INK
            cell.text_frame.word_wrap = True
    return shape


# ---------------------------------------------------------------- slides
# 1 — title
s = slide_base(1, "session", "workshop.start")
tb = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(12), Inches(2.6))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Как я работаю\n"
r.font.name = MONO; r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = INK
p2 = tf.add_paragraph()
r = p2.add_run(); r.text = "с "
r.font.name = MONO; r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = INK
r = p2.add_run(); r.text = "Claude Code"
r.font.name = MONO; r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = WARN
tb = s.shapes.add_textbox(Inches(0.7), Inches(5.05), Inches(10.5), Inches(1.1))
tf = tb.text_frame; tf.word_wrap = True
r = tf.paragraphs[0].add_run()
r.text = ("Контекст, скиллы, агенты и одна фича, которая вчера вечером "
          "прошла путь от спеки до двух pull request'ов.")
r.font.name = SANS; r.font.size = Pt(20); r.font.color.rgb = DIM
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.3), Inches(10), Inches(0.4))
r = tb.text_frame.paragraphs[0].add_run()
r.text = "Ruslan · tez.taxi mobile · 23.07.2026"
r.font.name = MONO; r.font.size = Pt(14); r.font.color.rgb = DIM

# 2 — plan
s = slide_base(2, "info", "workshop.plan")
title(s, "О чём поговорим", 36)
bullets(s, [
    ("Контекст", "— что модель видит, а что для неё не существует"),
    ("CLAUDE.md / AGENTS.md", "— инструкции, которые читаются всегда"),
    ("MD-файлы как память", "— worklog, хэндоффы, факты"),
    ("Скиллы", "— как повторяемый флоу становится плейбуком"),
    ("MCP против скилла", "— доступ к системе или знание процесса"),
    ("Плагины, хуки, агенты", "— что автоматизируем и кому делегируем"),
    ("Живой кейс", "— session file logs по SDD, со всеми граблями"),
], size=18, gap=10)

# 3 — context
s = slide_base(3, "warn", "basics.context")
title(s, "Контекст — это бюджет", 36, accent_words=("бюджет",))
paragraphs(s, [
    "У модели нет памяти. Есть окно контекста, порядка 200 тысяч токенов, и в него "
    "входит всё: системный промпт, CLAUDE.md, история диалога, вывод каждой команды. "
    "Чего нет в окне — того для модели не существует.",
    "Когда сессия разрастается, старое суммаризируется. Детали теряются, и модель "
    "начинает уверенно врать про то, что «помнит». Поэтому контекст я трачу как деньги: "
    "на задачу, а не на мусор.",
    "Пример экономии: хук rtk перехватывает git, grep и ls и отдаёт модели сжатый "
    "вывод. На длинной сессии это 60–90% сэкономленных токенов — буквально больше "
    "места для работы.",
])

# 4 — CLAUDE.md
s = slide_base(4, "info", "basics.claude_md")
title(s, "CLAUDE.md и AGENTS.md", 36)
paragraphs(s, [
    "Файлы, которые агент читает в начале каждой сессии. Глобальный ~/.claude/CLAUDE.md — "
    "личные правила. В репозитории — CLAUDE.md, импортирующий AGENTS.md: общий для "
    "Claude Code и Codex, оба агента живут по одним правилам.",
], top=1.95, size=17)
code_block(s, [
    ("# AGENTS.md — что туда попадает", DIM),
    ("toolchain: fvm, melos, Java 17          # нужно каждой сессии", INK),
    ("правила: коммиты только по команде      # защита от самодеятельности", INK),
    ("ключи 2GIS: проверка перед сборкой      # инцидент 2026-07-10", INK),
    ("роутинг: «для UI читай DESIGN.md»       # ссылка вместо простыни", INK),
], top=3.35, size=13)
paragraphs(s, [
    "Принцип — progressive disclosure: в always-loaded файле правила и ссылки, "
    "тяжёлые документы читаются только когда задача их касается.",
], top=5.55, size=17)

# 5 — md memory
s = slide_base(5, "info", "basics.md_memory")
title(s, "MD-файлы как память", 36)
table(s, [
    ("WORKLOG.md", "журнал сессий: что сделано, что осталось, грабли. Читается в начале, пополняется в конце"),
    ("HANDOFF-*.md", "передача задачи между сессиями: новая стартует с выводов, а не с археологии"),
    ("memory/*.md", "долговременные факты: один файл — один факт, плюс индекс MEMORY.md"),
    ("docs/adr/", "архитектурные решения с обоснованием, чтобы не спорить заново"),
], top=2.0, header=("файл", "зачем"))
paragraphs(s, ["Всё это — обычный markdown в git. Его читают три стороны: Claude, Codex и я."],
           top=5.6, size=17)

# 6 — skills what
s = slide_base(6, "info", "skills.what")
title(s, "Скилл — это плейбук с триггером", 34)
paragraphs(s, [
    "SKILL.md: описание-триггер модель видит всегда, тело с командами читает только "
    "при совпадении задачи. Сотня скиллов почти ничего не стоит, пока не понадобилась.",
], top=1.95, size=17)
table(s, [
    ("teztaxi-gitlab", "MR, пайплайны, мердж через glab"),
    ("run-app-ios", "запуск client/driver на симуляторе, с подменой ключей карт"),
    ("teztaxi-debug-harness", "репро багов: фейковый заказ, логгеры, adb/idb, жесты карты"),
    ("teztaxi-mr-review", "ревью чужого MR по чек-листу"),
    ("teztaxi-daily-report", "дневной отчёт «кто что делал»"),
], top=3.15, header=("мой скилл", "что автоматизирует"))

# 7 — skills how
s = slide_base(7, "info", "skills.how")
title(s, "Как рождается скилл", 36)
bullets(s, [
    ("", "Делаю флоу руками. Второй раз. Третий — и это уже сигнал"),
    ("", "Прошу Claude оформить прошедшую сессию в SKILL.md: он помнит команды и грабли лучше меня"),
    ("", "Правлю руками две вещи: триггер (когда срабатывать) и антипаттерны (чего не делать)"),
    ("", "Кладу в git. Наступили на новые грабли — дописали"),
], size=18, gap=12)
paragraphs(s, [
    "Хороший скилл: точный триггер, конкретные команды, явный список ошибок. "
    "Плохой скилл — статья. Хороший — инструкция по сборке.",
], top=5.2, size=17)

# 8 — mcp vs skill
s = slide_base(8, "warn", "skills.vs_mcp")
title(s, "MCP или скилл?", 36)
half = 5.9
for left, head, color, items in [
    (0.7, "MCP-сервер = доступ", INFO, [
        "инструменты и данные: живой процесс",
        "dcm — статический анализ Dart",
        "context7 — актуальные доки библиотек",
        "когда надо дёргать систему",
    ]),
    (6.9, "Скилл = знание", WARN, [
        "текст в git, ноль рантайма",
        "порядок действий, команды, чек-листы",
        "версионируется и ревьюится как код",
        "дешевле по токенам и поддержке",
    ]),
]:
    tb = s.shapes.add_textbox(Inches(left), Inches(2.0), Inches(half), Inches(3.4))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = head
    r.font.name = MONO; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = color
    p.space_after = Pt(10)
    for it in items:
        p = tf.add_paragraph(); p.space_after = Pt(7)
        r = p.add_run(); r.text = "> "
        r.font.name = MONO; r.font.size = Pt(16); r.font.color.rgb = color
        r = p.add_run(); r.text = it
        r.font.name = SANS; r.font.size = Pt(16); r.font.color.rgb = INK
paragraphs(s, [
    "Нужен доступ к API — MCP. Нужен порядок действий — скилл. "
    "Сомневаетесь — начните со скилла: чаще всего его хватает.",
], top=5.5, size=17)

# 9 — plugins
s = slide_base(9, "info", "plugins")
title(s, "Плагины: чужой опыт в коробке", 34)
table(s, [
    ("superpowers", "дисциплина процесса: brainstorming перед дизайном, TDD, планы. Не даёт кинуться писать код раньше спеки"),
    ("codex", "мост ко второй модели: Codex ревьюит спеки и код Claude"),
    ("humanizer", "вычищает ИИ-штампы: доки читаются как написанные человеком"),
    ("claude-mem", "автоматическая память и поиск по прошлым сессиям"),
    ("caveman", "сжатые ответы без воды, заодно экономит токены"),
], top=2.0, header=("плагин", "что даёт"))
paragraphs(s, ["Плагин — набор скиллов, агентов и хуков. Ставится один раз, работает во всех проектах."],
           top=5.85, size=17)

# 10 — hooks
s = slide_base(10, "warn", "hooks")
title(s, "Хуки: автоматика вне воли модели", 34)
paragraphs(s, [
    "Просьба «всегда делай так» живёт до первой суммаризации контекста. Хук — код, "
    "который срабатывает сам: модель не может его забыть или проигнорировать.",
], top=1.95, size=17)
code_block(s, [
    ("git status  → rtk git status    # тот же смысл, меньше токенов", INK),
    ("SessionStart → caveman on       # стиль ответов включён всегда", INK),
    ("PreToolUse   → проверка правил  # запреты, которые не обсуждаются", INK),
], top=3.3, size=14)
paragraphs(s, ["Поведение должно быть гарантированным — значит это хук, а не строчка в промпте."],
           top=5.2, size=17)

# 11 — agents
s = slide_base(11, "info", "agents")
title(s, "Агенты: свежая голова под задачу", 34)
paragraphs(s, [
    "Субагент — отдельное окно контекста. Главный тред держит задачу целиком, субагенты "
    "делают грязную работу и возвращают короткий отчёт, а не тонну прочитанных файлов.",
], top=1.9, size=17)
bullets(s, [
    ("Explore", "— прочесать кодовую базу; моё окно остаётся чистым"),
    ("Plan", "— архитектурный план до того, как трогаем код"),
    ("Доменные агенты", "— по фичам приложения, каждый знает свой кусок"),
    ("Параллельность", "— три независимых поиска бегут одновременно"),
], top=3.1, size=17, gap=8)
paragraphs(s, [
    "Правило оркестрации: делегируй вывод, а не решение. Решения принимает тред, "
    "у которого весь контекст задачи.",
], top=5.5, size=17)

# 12 — cross-model
s = slide_base(12, "error", "agents.cross_model")
title(s, "Вторая модель как ревьюер", 34)
paragraphs(s, [
    "Вчера Codex ревьюил спеку и код фичи, которую писал Claude. Результат: два P0 и восемь P1.",
], top=1.95, size=17)
bullets(s, [
    ("P0:", "flush() не был безопасным барьером — гонка над RandomAccessFile, который запрещает параллельные вызовы"),
    ("P0:", "выгрузка отдавала живые файлы: ротация могла удалить файл во время аплоада"),
    ("P1:", "после чанка p99 файлы не парсились; циклический data отключал storage; dart:io в главном entrypoint ломал web"),
], top=2.9, size=16, gap=10)
paragraphs(s, [
    "У разных моделей разные слепые зоны. Перекрёстное ревью стоит десять минут "
    "и дешевле любого продакшн-инцидента.",
], top=5.4, size=17)

# 13 — handoffs
s = slide_base(13, "warn", "handoffs")
title(s, "Хэндоффы: сессия конечна", 36)
paragraphs(s, [
    "Любая сессия упирается в лимит контекста. Заканчиваю большую работу — прошу агента "
    "написать HANDOFF: что сделано, что осталось, какие грабли уже найдены. Следующая "
    "сессия читает три абзаца и продолжает с места.",
    "То же между агентами: Claude оставляет хэндофф, Codex подхватывает. Формат один — "
    "markdown в репозитории.",
    "Главное правило: переносим выводы, не переписку. Хэндофф на страницу полезнее "
    "лога на мегабайт.",
])

# 14 — case timeline
s = slide_base(14, "case", "case.sdd")
title(s, "Кейс: session file logs за вечер", 32)
paragraphs(s, [
    "Логи каждой сессии приложения хранятся на устройстве, переживают краш, юзер "
    "отправляет их в саппорт кнопкой. Работали по SDD — сначала спека, потом код.",
], top=1.85, size=16)
code_block(s, [
    ("18:05  спека: лимиты, ротация, pitfalls        brainstorming → docs/specs/", INK),
    ("18:20  Codex ревьюит спеку                     параллельно пишется план", INK),
    ("18:30  TDD: options → naming → encoder → core  тест первым, потом код", INK),
    ("19:00  фидбек: «maxSessions не нужен,          спека правится дешевле кода", WARN),
    ("       выгрузку — добавить»", WARN),
    ("19:40  Codex: 2×P0, 8×P1                       коммит-хардеринг", INK),
    ("20:10  zip-экспорт, share_plus, доки           humanizer для текстов", INK),
    ("20:30  два PR: team_logger + flutter_team_logger", OK),
], top=3.0, size=13)

# 15 — case decisions
s = slide_base(15, "case", "case.decisions")
title(s, "Что решали в спеке", 34)
table(s, [
    ("сколько живут", "7 дней + потолок каталога 20 MiB. Лимита на число сессий нет: возраст и вес ограничивают сами"),
    ("сессия растёт", "чанки по 512 KiB; переполнилось — гибнет старейший чанк. Хвост важнее начала"),
    ("запись-гигант", "32 KiB на запись; режем stack, потом data, сообщение — последним"),
    ("краш при записи", "JSON Lines: страдает максимум последняя строка"),
    ("диск полон", "5 неудач подряд — storage выключает себя; исключений в приложение нет"),
    ("отправка", "zip-снапшот, неизменяемый: безопасно аплоадить при живой записи"),
], top=1.95, size=13, header=("вопрос", "решение"))

# 16 — lessons
s = slide_base(16, "info", "case.lessons")
title(s, "Что из этого выношу", 36)
bullets(s, [
    ("", "Спека на две страницы читается за пять минут — и правится за пять минут. Код правится за час"),
    ("", "Человек и вторая модель ловят разное: человек — «зачем этот лимит», Codex — гонку над дескриптором"),
    ("", "Тест первым — не ритуал: красный тест дважды ловил мои ошибки до коммита"),
    ("", "Атомарные коммиты дают дешёвый откат любого решения"),
    ("", "Спека и план в git рядом с кодом: ревьюер видит не только «что», но и «почему»"),
], size=17, gap=11)

# 17 — start
s = slide_base(17, "ok", "start")
title(s, "С чего начать", 36)
bullets(s, [
    ("1. CLAUDE.md в репозитории.", "Toolchain, команды, правила. Час работы, окупается в первый день"),
    ("2. Один скилл", "на самую частую рутину: создание MR или запуск на симуляторе"),
    ("3. WORKLOG.md.", "Пять строк в конце сессии, чтобы завтра не начинать с нуля"),
    ("4. Ревью второй моделью", "для крупных фич — когда первые три пункта приживутся"),
], size=18, gap=12)
paragraphs(s, [
    "Не пытайтесь завести всё сразу. Каждый инструмент появлялся после конкретной боли, "
    "а не из каталога.",
], top=5.3, size=17)

# 18 — QA
s = slide_base(18, "session_end", "workshop.qa")
tb = s.shapes.add_textbox(Inches(0.7), Inches(2.3), Inches(12), Inches(1.2))
r = tb.text_frame.paragraphs[0].add_run()
r.text = "Вопросы"
r.font.name = MONO; r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = INK
bullets(s, [
    ("ядро:", "github.com/Ruslan660/team_logger — PR #1"),
    ("Flutter-обвязка и спека:", "github.com/Ruslan660/flutter_team_logger — PR #1"),
    ("спека и план:", "docs/specs/ и docs/plans/ в репозитории воркшопа"),
], top=4.0, size=17, gap=9)

out = "/Users/ucantrustmedude/Learn/ai-workshop-teztaxi/docs/slides/ai-workshop.pptx"
prs.save(out)
print("saved", out)
